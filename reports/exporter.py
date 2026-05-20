"""导出 JSON / HTML / Excel / 文本巡检报告（企业交付）。"""
from __future__ import annotations

import html as html_lib
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

from core.sensor_display import (
    NO_REPORT_MSG,
    OFF_SENSOR_MSG,
    apply_sensor_display,
    format_sensor_value,
    is_placeholder_display,
)

try:
    from jinja2 import Template as _JinjaTemplate
except ImportError:
    _JinjaTemplate = None

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    _HAS_XLSX = True
except ImportError:
    _HAS_XLSX = False

_HTML_STYLE = """
body{font-family:Segoe UI,Microsoft YaHei,sans-serif;margin:24px;background:#fafafa;}
h1{color:#0d47a1;} .meta{color:#555;} table{border-collapse:collapse;width:100%;margin:12px 0;}
th,td{border:1px solid #ccc;padding:8px;text-align:left;font-size:13px;}
th{background:#e3f2fd;} .err{background:#ffebee;color:#c62828;font-weight:bold;}
.warn{background:#fff8e1;color:#ef6c00;} .ok{background:#e8f5e9;}
pre{background:#fff;border:1px solid #ddd;padding:12px;overflow:auto;}
"""


def _timestamp() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def export_json(report_dir: Path, data: Dict[str, Any], prefix: str = "report") -> Path:
    report_dir.mkdir(parents=True, exist_ok=True)
    out = report_dir / f"{prefix}_{_timestamp()}.json"
    payload = dict(data)
    if payload.get("inspection"):
        payload["inspection"] = apply_sensor_display(payload["inspection"])
    out.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return out


def _simple_render(template: str, ctx: Dict[str, Any]) -> str:
    out = template

    def esc(v: Any) -> str:
        return html_lib.escape(str(v))

    for m in re.finditer(r"\{\%\s*for\s+(\w+)\s+in\s+(\w+)\s*\%\}(.*?)\{\%\s*endfor\s*\%\}", out, re.S):
        var, seq_name, block = m.group(1), m.group(2), m.group(3)
        seq: List[Any] = ctx.get(seq_name) or []
        parts = []
        for item in seq:
            row = block
            if isinstance(item, dict):
                for k, v in item.items():
                    row = row.replace("{{ " + k + " }}", esc(v))
                    row = row.replace("{{" + k + "}}", esc(v))
            else:
                row = row.replace("{{ " + var + " }}", esc(item))
            parts.append(row)
        out = out[: m.start()] + "".join(parts) + out[m.end() :]

    for key, val in ctx.items():
        if isinstance(val, list):
            out = out.replace("{{ " + key + "|length }}", str(len(val)))
            out = out.replace("{{" + key + "|length}}", str(len(val)))
        else:
            out = out.replace("{{ " + key + " }}", esc(val))
            out = out.replace("{{" + key + "}}", esc(val))
    return out


HTML_TEMPLATE_SRC = """
<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8"><title>BMC AutoInsight 巡检报告</title>
<style>""" + _HTML_STYLE + """</style>
</head>
<body>
<h1>BMC AutoInsight 巡检报告</h1>
<p class="meta">目标: {{ host }} | Mock: {{ mock }} | 时间: {{ ts }}</p>
<h2>异常与说明</h2>
<table>
<tr><th>类型</th><th>说明</th></tr>
{% for n in notes %}<tr class="warn"><td>提示</td><td>{{ n }}</td></tr>{% endfor %}
{% for f in faults %}<tr class="err"><td>{{ f.severity }}</td><td>{{ f.rule }}: {{ f.message }}</td></tr>{% endfor %}
</table>
<h2>存储/事件日志状态</h2>
<table>
<tr><th>模块</th><th>状态说明</th></tr>
<tr><td>Storage</td><td>{{ storage_notice }}</td></tr>
<tr><td>EventLog</td><td>{{ event_notice }}</td></tr>
</table>
<h2>温度 / 风扇 / 功耗</h2>
<table><tr><th>类型</th><th>名称</th><th>读数</th></tr>
{% for t in thermal_rows %}<tr><td>温度</td><td>{{ t.name }}</td><td>{{ t.val }}</td></tr>{% endfor %}
{% for f in fan_rows %}<tr><td>风扇</td><td>{{ f.name }}</td><td>{{ f.val }}</td></tr>{% endfor %}
{% for p in power_rows %}<tr><td>功耗</td><td>{{ p.name }}</td><td>{{ p.val }}</td></tr>{% endfor %}
</table>
<h2>硬件信息</h2>
<table><tr><th>项</th><th>值</th></tr>
{% for h in hardware_summary_rows %}<tr><td>{{ h.label }}</td><td>{{ h.value }}</td></tr>{% endfor %}
</table>
<h3>主板</h3>
<table><tr><th>项</th><th>值</th></tr>
{% for m in motherboard_rows %}<tr><td>{{ m.label }}</td><td>{{ m.value }}</td></tr>{% endfor %}
</table>
<h3>CPU</h3>
<table><tr><th>ID</th><th>型号</th><th>核</th><th>线程</th><th>频率MHz</th><th>厂商</th></tr>
{% for c in cpu_rows %}<tr><td>{{ c.id }}</td><td>{{ c.model }}</td><td>{{ c.cores }}</td><td>{{ c.threads }}</td><td>{{ c.speed_mhz }}</td><td>{{ c.manufacturer }}</td></tr>{% endfor %}
</table>
<h3>内存</h3>
<p>总量: {{ memory_total }} GiB</p>
<table><tr><th>ID</th><th>容量GiB</th><th>类型</th><th>频率MHz</th><th>厂商</th><th>序列号</th></tr>
{% for m in memory_rows %}<tr><td>{{ m.id }}</td><td>{{ m.capacity_gib }}</td><td>{{ m.type }}</td><td>{{ m.speed_mhz }}</td><td>{{ m.manufacturer }}</td><td>{{ m.serial }}</td></tr>{% endfor %}
</table>
<h3>硬盘</h3>
<table><tr><th>ID</th><th>型号</th><th>容量GiB</th><th>协议</th><th>介质</th><th>健康</th></tr>
{% for d in disk_rows %}<tr><td>{{ d.id }}</td><td>{{ d.model }}</td><td>{{ d.capacity_gib }}</td><td>{{ d.protocol }}</td><td>{{ d.media_type }}</td><td>{{ d.health }}</td></tr>{% endfor %}
</table>
<h2>系统摘要</h2>
<pre>{{ system_json }}</pre>
<h2>API 冒烟</h2>
<p>总计 {{ api_total }} | 通过 {{ api_passed }} | 失败 {{ api_failed }}</p>
<table><tr><th>用例</th><th>路径</th><th>HTTP</th><th>耗时ms</th><th>结果</th></tr>
{% for a in api_results %}<tr class="{{ a.row_class }}">
<td>{{ a.name }}</td><td>{{ a.path }}</td><td>{{ a.status }}</td><td>{{ a.duration_ms }}</td><td>{{ a.passed_text }}</td></tr>{% endfor %}
</table>
</body></html>
"""


def _render_html(ctx: Dict[str, Any]) -> str:
    if _JinjaTemplate is not None:
        return _JinjaTemplate(HTML_TEMPLATE_SRC).render(**ctx)
    try:
        from reports.jinja2_inline import Template as _InlineTemplate

        return _InlineTemplate(HTML_TEMPLATE_SRC).render(**ctx)
    except Exception:
        return _simple_render(HTML_TEMPLATE_SRC, ctx)


def _kv_rows(mapping: Dict[str, Any], labels: Dict[str, str]) -> List[Dict[str, str]]:
    return [{"label": labels.get(k, k), "value": mapping.get(k) or "-"} for k in labels]


def _prepare_hardware_ctx(hw: Dict[str, Any]) -> Dict[str, Any]:
    hw = hw or {}
    summary = hw.get("summary") or {}
    mb = hw.get("motherboard") or {}
    mem = hw.get("memory") or {}
    return {
        "hardware_summary_rows": _kv_rows(
            summary,
            {
                "manufacturer": "制造商",
                "model": "型号",
                "serial_number": "序列号",
                "host_name": "主机名",
                "power_state": "PowerState",
                "bios_version": "BIOS 版本",
                "bmc_firmware_version": "BMC 固件",
            },
        ),
        "motherboard_rows": _kv_rows(
            mb,
            {
                "manufacturer": "制造商",
                "model": "型号",
                "part_number": "部件号",
                "serial_number": "序列号",
                "sku": "SKU",
            },
        ),
        "cpu_rows": hw.get("cpu") or [{"id": "-", "model": "(无)", "cores": "-", "threads": "-", "speed_mhz": "-", "manufacturer": "-"}],
        "memory_total": mem.get("total_gib") if mem.get("total_gib") is not None else "-",
        "memory_rows": mem.get("modules")
        or [{"id": "-", "capacity_gib": "-", "type": "-", "speed_mhz": "-", "manufacturer": "-", "serial": "-"}],
        "disk_rows": hw.get("disks")
        or [{"id": "-", "model": hw.get("storage_notice") or "(无硬盘详情)", "capacity_gib": "-", "protocol": "-", "media_type": "-", "health": "-"}],
    }


def _prepare_api_rows(raw: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []
    for a in raw or []:
        passed = bool(a.get("passed"))
        rows.append(
            {
                "name": a.get("name") or "-",
                "path": a.get("path") or "-",
                "status": a.get("status") if a.get("status") is not None else "ERR",
                "duration_ms": a.get("duration_ms", ""),
                "passed_text": "PASS" if passed else "FAIL",
                "row_class": "ok" if passed else "err",
            }
        )
    if not rows:
        rows.append(
            {
                "name": "(未执行冒烟测试)",
                "path": "-",
                "status": "-",
                "duration_ms": "-",
                "passed_text": "-",
                "row_class": "warn",
            }
        )
    return rows


def export_html(report_dir: Path, data: Dict[str, Any], prefix: str = "inspect") -> Path:
    report_dir.mkdir(parents=True, exist_ok=True)
    out = report_dir / f"{prefix}_{_timestamp()}.html"
    inv_raw = data.get("inspection")
    if isinstance(inv_raw, dict) and inv_raw:
        inv = apply_sensor_display(inv_raw)
    else:
        inv = {}
    faults = data.get("faults")
    if faults is None and isinstance(inv, dict):
        faults = inv.get("faults")
    if faults is None:
        faults = []
    notes = inv.get("inspection_notes") or []
    storage = inv.get("storage") or {}
    event = inv.get("event_log") or {}
    thermal = inv.get("thermal") or {}
    power = inv.get("power") or {}
    ps = inv.get("power_state") or (inv.get("system") or {}).get("PowerState")

    thermal_rows = [
        {"name": t.get("Name") or t.get("MemberId"), "val": format_sensor_value(t.get("ReadingCelsius"), ps)}
        for t in thermal.get("Temperatures") or []
    ]
    fan_rows = [
        {
            "name": f.get("Name") or f.get("MemberId"),
            "val": format_sensor_value(f.get("Reading") or f.get("ReadingRPM"), ps),
        }
        for f in thermal.get("Fans") or []
    ]
    power_rows = [
        {"name": pc.get("Name") or pc.get("MemberId"), "val": format_sensor_value(pc.get("PowerConsumedWatts"), ps)}
        for pc in power.get("PowerControl") or []
    ]

    raw_api = data.get("api_test") or inv.get("api_results") or []
    api_results = _prepare_api_rows(raw_api)
    hw = data.get("hardware_info") or inv.get("hardware_info") or {}
    hw_ctx = _prepare_hardware_ctx(hw)

    ctx = {
        "host": inv.get("host") or data.get("host"),
        "mock": inv.get("mock") if inv.get("mock") is not None else data.get("mock"),
        "ts": inv.get("timestamp") or data.get("timestamp"),
        "faults": faults,
        "notes": notes,
        "storage_notice": storage.get("customer_notice_zh") or "正常或未跳过",
        "event_notice": event.get("customer_notice_zh")
        or event.get("_event_log_probe_summary")
        or "正常或未跳过",
        "system_json": json.dumps(inv.get("system") or {}, ensure_ascii=False, indent=2),
        "api_results": api_results,
        "api_total": len(raw_api),
        "api_passed": sum(1 for x in raw_api if x.get("passed")),
        "api_failed": sum(1 for x in raw_api if not x.get("passed")),
        "thermal_rows": thermal_rows,
        "fan_rows": fan_rows,
        "power_rows": power_rows,
        **hw_ctx,
    }
    out.write_text(_render_html(ctx), encoding="utf-8")
    return out


def export_txt(report_dir: Path, data: Dict[str, Any], prefix: str = "inspect") -> Path:
    report_dir.mkdir(parents=True, exist_ok=True)
    out = report_dir / f"{prefix}_{_timestamp()}.txt"
    inv_raw = data.get("inspection")
    inv = apply_sensor_display(inv_raw) if isinstance(inv_raw, dict) and inv_raw else {}
    lines = [
        "BMC AutoInsight 巡检报告（文本）",
        f"时间: {data.get('timestamp') or inv.get('timestamp')}",
        f"目标: {data.get('host') or inv.get('host')}",
        "",
        "【提示与跳过说明】",
    ]
    for n in inv.get("inspection_notes") or []:
        lines.append(f" - {n}")
    st = inv.get("storage") or {}
    if st.get("customer_notice_zh"):
        lines.append(f"存储: {st['customer_notice_zh']}")
    ev = inv.get("event_log") or {}
    if ev.get("customer_notice_zh") or ev.get("_event_log_probe_summary"):
        lines.append(f"事件: {ev.get('customer_notice_zh') or ev.get('_event_log_probe_summary')}")
    lines.append("")
    lines.append("【传感器】")
    ps = inv.get("power_state") or (inv.get("system") or {}).get("PowerState")
    for t in (inv.get("thermal") or {}).get("Temperatures") or []:
        name = t.get("Name") or t.get("MemberId")
        val = format_sensor_value(t.get("ReadingCelsius"), ps)
        lines.append(f"  温度 {name}: {val}")
    for f in (inv.get("thermal") or {}).get("Fans") or []:
        name = f.get("Name") or f.get("MemberId")
        val = format_sensor_value(f.get("Reading") or f.get("ReadingRPM"), ps)
        lines.append(f"  风扇 {name}: {val}")
    for pc in (inv.get("power") or {}).get("PowerControl") or []:
        name = pc.get("Name") or pc.get("MemberId")
        val = format_sensor_value(pc.get("PowerConsumedWatts"), ps)
        lines.append(f"  功耗 {name}: {val}")
    lines.append("")
    lines.append("【API 冒烟】")
    raw_api = data.get("api_test") or inv.get("api_results") or []
    if not raw_api:
        lines.append(" (未执行)")
    for a in raw_api:
        ok = "PASS" if a.get("passed") else "FAIL"
        lines.append(f"  {a.get('name')}: {a.get('path')} HTTP={a.get('status')} {ok}")
    lines.append("")
    hw = data.get("hardware_info") or inv.get("hardware_info") or {}
    if hw:
        s = hw.get("summary") or {}
        lines.append("【硬件信息】")
        lines.append(f"  型号: {s.get('model')}  SN: {s.get('serial_number')}")
        lines.append(f"  BIOS: {s.get('bios_version')}  BMC: {s.get('bmc_firmware_version')}")
        lines.append("")
    lines.append("【告警项】")
    for f in data.get("faults") or inv.get("faults") or []:
        lines.append(f" - [{f.get('severity')}] {f.get('rule')}: {f.get('message')}")
    lines.append("")
    out.write_text("\n".join(lines), encoding="utf-8")
    return out


def export_excel(report_dir: Path, data: Dict[str, Any], prefix: str = "inspect") -> Optional[Path]:
    """导出 Excel；若未安装 openpyxl 则返回 None。"""
    if not _HAS_XLSX:
        return None
    report_dir.mkdir(parents=True, exist_ok=True)
    out = report_dir / f"{prefix}_{_timestamp()}.xlsx"

    if data.get("batch"):
        wb = Workbook()
        ws = wb.active
        ws.title = "Batch"
        ws.append(["主机", "告警数", "错误"])
        for item in data.get("results") or []:
            ws.append(
                [
                    item.get("host"),
                    len(item.get("faults") or []),
                    item.get("error", ""),
                ]
            )
        wb.save(out)
        return out

    inv_raw = data.get("inspection")
    inv = apply_sensor_display(inv_raw) if isinstance(inv_raw, dict) and inv_raw else {}
    wb = Workbook()
    ps = inv.get("power_state") or (inv.get("system") or {}).get("PowerState")

    ws0 = wb.active
    ws0.title = "Summary"
    ws0.append(["字段", "值"])
    ws0.append(["主机", data.get("host") or inv.get("host")])
    ws0.append(["时间", data.get("timestamp") or inv.get("timestamp")])
    ws0.append(["PowerState", ps])
    ws0.append(["Mock", str(inv.get("mock"))])
    for n in inv.get("inspection_notes") or []:
        ws0.append(["提示", n])

    ws1 = wb.create_sheet("Faults")
    ws1.append(["级别", "规则", "消息"])
    red_font = Font(color="C62828", bold=True)
    for f in data.get("faults") or inv.get("faults") or []:
        ws1.append([f.get("severity"), f.get("rule"), f.get("message")])
        for cell in ws1[ws1.max_row]:
            cell.font = red_font

    ws2 = wb.create_sheet("Thermal")
    ws2.append(["传感器", "读数°C", "健康"])
    off_fill = PatternFill("solid", fgColor="FFEBEE")
    no_data_fill = PatternFill("solid", fgColor="FFF9C4")

    def _style_sensor_row(ws, val: Any) -> None:
        if val == OFF_SENSOR_MSG:
            fill = off_fill
        elif val == NO_REPORT_MSG:
            fill = no_data_fill
        else:
            return
        for cell in ws[ws.max_row]:
            cell.fill = fill

    for t in (inv.get("thermal") or {}).get("Temperatures") or []:
        val = t.get("ReadingCelsius")
        ws2.append([t.get("Name") or t.get("MemberId"), val, str((t.get("Status") or {}).get("Health"))])
        _style_sensor_row(ws2, val)

    ws3 = wb.create_sheet("Fans")
    ws3.append(["风扇", "转速", "健康"])
    for fn in (inv.get("thermal") or {}).get("Fans") or []:
        val = fn.get("Reading") or fn.get("ReadingRPM")
        ws3.append([fn.get("Name") or fn.get("MemberId"), val, str((fn.get("Status") or {}).get("Health"))])
        _style_sensor_row(ws3, val)

    ws_pwr = wb.create_sheet("Power")
    ws_pwr.append(["项", "读数"])
    for pc in (inv.get("power") or {}).get("PowerControl") or []:
        val = pc.get("PowerConsumedWatts")
        ws_pwr.append([pc.get("Name") or pc.get("MemberId"), val])
        _style_sensor_row(ws_pwr, val)

    ws4 = wb.create_sheet("Modules")
    ws4.append(["模块", "说明"])
    st = inv.get("storage") or {}
    ws4.append(["Storage", st.get("customer_notice_zh") or ""])
    ev = inv.get("event_log") or {}
    ws4.append(["EventLog", ev.get("customer_notice_zh") or ev.get("_event_log_probe_summary") or ""])

    hw = data.get("hardware_info") or inv.get("hardware_info") or {}
    if hw:
        ws_hw = wb.create_sheet("Hardware")
        ws_hw.append(["项", "值"])
        summary = hw.get("summary") or {}
        for label, key in (
            ("制造商", "manufacturer"),
            ("型号", "model"),
            ("序列号", "serial_number"),
            ("BIOS", "bios_version"),
            ("BMC固件", "bmc_firmware_version"),
            ("PowerState", "power_state"),
        ):
            ws_hw.append([label, summary.get(key) or ""])
        mb = hw.get("motherboard") or {}
        ws_hw.append(["主板型号", mb.get("model") or ""])
        ws_hw.append(["主板部件号", mb.get("part_number") or ""])
        ws_hw.append(["主板序列号", mb.get("serial_number") or ""])

        ws_cpu = wb.create_sheet("CPU")
        ws_cpu.append(["ID", "型号", "核", "线程", "MHz", "厂商"])
        for c in hw.get("cpu") or []:
            ws_cpu.append(
                [c.get("id"), c.get("model"), c.get("cores"), c.get("threads"), c.get("speed_mhz"), c.get("manufacturer")]
            )

        ws_mem = wb.create_sheet("Memory")
        mem = hw.get("memory") or {}
        ws_mem.append(["总容量GiB", mem.get("total_gib") or ""])
        ws_mem.append(["ID", "容量GiB", "类型", "MHz", "厂商", "序列号"])
        for m in mem.get("modules") or []:
            ws_mem.append(
                [m.get("id"), m.get("capacity_gib"), m.get("type"), m.get("speed_mhz"), m.get("manufacturer"), m.get("serial")]
            )

        ws_disk = wb.create_sheet("Disks")
        ws_disk.append(["ID", "型号", "容量GiB", "协议", "介质", "健康"])
        for d in hw.get("disks") or []:
            ws_disk.append(
                [d.get("id"), d.get("model"), d.get("capacity_gib"), d.get("protocol"), d.get("media_type"), d.get("health")]
            )

    raw_api = data.get("api_test") or inv.get("api_results") or []
    ws_api = wb.create_sheet("API_Smoke")
    ws_api.append(["用例", "路径", "HTTP", "耗时ms", "结果"])
    api_rows = _prepare_api_rows(raw_api)
    for a in api_rows:
        ws_api.append([a.get("name"), a.get("path"), a.get("status"), a.get("duration_ms"), a.get("passed_text")])
        if a.get("row_class") == "err":
            for cell in ws_api[ws_api.max_row]:
                cell.font = red_font

    warn_fill = PatternFill("solid", fgColor="FFF9C4")
    for row in ws4.iter_rows(min_row=2, max_row=ws4.max_row):
        for c in row:
            if c.value and "权限" in str(c.value):
                c.fill = warn_fill

    wb.save(out)
    return out


def export_inspection_bundle(
    report_dir: Path,
    payload: Dict[str, Any],
    exports: str = "json,html,txt,xlsx",
    prefix: str = "inspect",
) -> Dict[str, Optional[str]]:
    """一次导出多种格式；exports 为逗号分隔：json,html,txt,xlsx。"""
    want = {x.strip().lower() for x in exports.replace(";", ",").split(",") if x.strip()}
    paths: Dict[str, Optional[str]] = {}
    if "json" in want:
        p = export_json(report_dir, payload, prefix=prefix)
        paths["json"] = str(p)
    if "html" in want:
        paths["html"] = str(export_html(report_dir, payload, prefix=prefix))
    if "txt" in want:
        paths["txt"] = str(export_txt(report_dir, payload, prefix=prefix))
    if "xlsx" in want or "excel" in want:
        px = export_excel(report_dir, payload, prefix=prefix)
        paths["xlsx"] = str(px) if px else "(未安装 openpyxl，跳过 Excel)"
    return paths
